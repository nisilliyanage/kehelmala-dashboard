"""
Step 1 of the pipeline.

Reads every export dropped into data/raw/{youtube,facebook,instagram,tiktok}/
and writes a single flat list of posts to data/processed/all_posts.json.

Each post dict has: platform, title, permalink, views, likes, comments, shares,
date (when posted), asof (date the export itself is "as of" — comes ONLY from the
export's own data, never from today's date), post_id (for de-duplication), post_type.
"""
import csv
import glob
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config

RAW_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'raw')
OUT_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'processed', 'all_posts.json')

EXPORT_DATE_RE = re.compile(r'([A-Za-z]+-\d{1,2}-\d{4})_([A-Za-z]+-\d{1,2}-\d{4})')


def _asof_from_filename(filename):
    """Meta exports encode their date range in the filename, e.g.
    Apr-15-2026_Sep-28-2026_....csv -> asof = 2026-09-28. Returns None if no match,
    so callers never fall back to today's date."""
    m = EXPORT_DATE_RE.search(os.path.basename(filename))
    if not m:
        return None
    end = m.group(2)
    try:
        import datetime
        dt = datetime.datetime.strptime(end, '%b-%d-%Y')
        return dt.strftime('%Y-%m-%d')
    except ValueError:
        return None


def load_youtube():
    """Reads the Videos + Metrics sheets of the combined xlsx export (built by the
    legacy yt_pull.py / yt_merge_shares.py scripts). asof comes ONLY from the xlsx's
    own 'Data as of' column; None if that column is blank."""
    xlsx_files = glob.glob(os.path.join(RAW_DIR, 'youtube', '*.xlsx'))
    if not xlsx_files:
        return load_youtube_from_studio_csv_only()

    # Pre-load durations and extra info from Studio CSV(s) if present
    durations_by_id = {}
    durations_by_title = {}
    csv_shares_by_id = {}
    for csv_path in glob.glob(os.path.join(RAW_DIR, 'youtube', '*.csv')):
        with open(csv_path, newline='', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                cid = (row.get('Content') or '').strip()
                title = (row.get('Video title') or '').strip()
                dur = row.get('Duration')
                shares_val = row.get('Shares')
                if cid and cid != 'Total':
                    if dur:
                        try:
                            durations_by_id[cid] = int(float(dur))
                        except ValueError:
                            pass
                    if shares_val:
                        try:
                            csv_shares_by_id[cid] = int(float(str(shares_val).replace(',', '')))
                        except ValueError:
                            pass
                if title and dur and cid != 'Total':
                    try:
                        durations_by_title[title] = int(float(dur))
                    except ValueError:
                        pass

    import datetime
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_files[0], data_only=True)
    posts = []

    if 'Metrics' in wb.sheetnames and 'Videos' in wb.sheetnames:
        m_ws = wb['Metrics']
        v_ws = wb['Videos']

        v_headers = [str(c.value).strip() if c.value else '' for c in v_ws[1]]
        def v_col(names):
            for n in names:
                if n in v_headers:
                    return v_headers.index(n)
            return None
        v_idx_title = v_col(['Video', 'Title', 'Video title'])
        v_idx_date = v_col(['Release date', 'Video publish time', 'Date'])
        v_idx_url = v_col(['Notes', 'Video URL', 'URL', 'Link'])

        v_data_by_row = {}
        v_data_by_title = {}
        for r in range(2, v_ws.max_row + 1):
            t_val = v_ws.cell(row=r, column=v_idx_title + 1).value if v_idx_title is not None else ''
            d_val = v_ws.cell(row=r, column=v_idx_date + 1).value if v_idx_date is not None else None
            u_val = v_ws.cell(row=r, column=v_idx_url + 1).value if v_idx_url is not None else ''

            t_str = str(t_val).strip() if t_val else ''
            if isinstance(d_val, (datetime.datetime, datetime.date)):
                d_str = d_val.strftime('%Y-%m-%d')
            else:
                d_str = str(d_val)[:10] if d_val else None
            u_str = str(u_val).strip() if u_val else ''

            info = {'title': t_str, 'date': d_str, 'url': u_str}
            v_data_by_row[r] = info
            if t_str:
                v_data_by_title[t_str] = info

        m_headers = [str(c.value).strip() if c.value else '' for c in m_ws[1]]
        def m_col(names):
            for n in names:
                if n in m_headers:
                    return m_headers.index(n)
            return None
        idx_title = m_col(['Video', 'Title', 'Video title'])
        idx_views = m_col(['Views', 'Total views'])
        idx_likes = m_col(['Reactions/Likes', 'Likes', 'Reactions'])
        idx_shares = m_col(['Shares/Reposts', 'Shares', 'Total shares'])
        idx_comments = m_col(['Comments', 'Total comments'])
        idx_asof = m_col(['Data as of', 'As of'])

        for r in range(2, m_ws.max_row + 1):
            title = m_ws.cell(row=r, column=idx_title + 1).value if idx_title is not None else None
            if not title:
                continue
            title_str = str(title).strip()
            v_info = v_data_by_row.get(r) or v_data_by_title.get(title_str) or {}
            url = v_info.get('url', '')
            vid = url.split('v=')[-1] if 'v=' in url else ''
            duration = durations_by_id.get(vid) or durations_by_title.get(title_str)

            asof_val = m_ws.cell(row=r, column=idx_asof + 1).value if idx_asof is not None else None
            if isinstance(asof_val, (datetime.datetime, datetime.date)):
                asof = asof_val.strftime('%Y-%m-%d')
            else:
                asof = str(asof_val)[:10] if asof_val else None

            views = int(m_ws.cell(row=r, column=idx_views + 1).value or 0) if idx_views is not None else 0
            likes = int(m_ws.cell(row=r, column=idx_likes + 1).value or 0) if idx_likes is not None else 0
            comments = int(m_ws.cell(row=r, column=idx_comments + 1).value or 0) if idx_comments is not None else 0

            raw_shares = m_ws.cell(row=r, column=idx_shares + 1).value if idx_shares is not None else None
            shares = int(raw_shares or 0) if raw_shares is not None else csv_shares_by_id.get(vid, 0)

            posts.append({
                'platform': 'YouTube',
                'title': title_str,
                'permalink': url,
                'views': views,
                'likes': likes,
                'comments': comments,
                'shares': shares,
                'date': v_info.get('date'),
                'asof': asof,
                'duration_sec': duration,
                'post_id': url or title_str,
                'post_type': 'video',
            })
    else:
        # Fallback for single-sheet workbook
        ws = wb.active
        headers = [str(c.value).strip() if c.value else '' for c in ws[1]]
        def col(names):
            for n in names:
                if n in headers:
                    return headers.index(n)
            return None
        idx_title = col(['Video title', 'Title', 'Video', 'Content'])
        idx_url = col(['Video URL', 'URL', 'Notes', 'Link', 'Permalink'])
        idx_views = col(['Views', 'Total views'])
        idx_likes = col(['Reactions/Likes', 'Likes', 'Reactions'])
        idx_comments = col(['Comments', 'Total comments'])
        idx_shares = col(['Shares/Reposts', 'Shares', 'Total shares'])
        idx_date = col(['Video publish time', 'Release date', 'Publish time', 'Date'])
        idx_asof = col(['Data as of', 'As of'])
        idx_duration = col(['Duration'])

        for row in ws.iter_rows(min_row=2, values_only=True):
            if idx_title is None or row[idx_title] in (None, ''):
                continue
            asof = None
            if idx_asof is not None and row[idx_asof]:
                asof_raw = row[idx_asof]
                if isinstance(asof_raw, (datetime.datetime, datetime.date)):
                    asof = asof_raw.strftime('%Y-%m-%d')
                else:
                    asof = str(asof_raw)[:10]

            date_val = None
            if idx_date is not None and row[idx_date]:
                raw_d = row[idx_date]
                if isinstance(raw_d, (datetime.datetime, datetime.date)):
                    date_val = raw_d.strftime('%Y-%m-%d')
                else:
                    date_val = str(raw_d)[:10]

            u_str = str(row[idx_url]) if idx_url is not None and row[idx_url] else ''
            vid = u_str.split('v=')[-1] if 'v=' in u_str else ''
            dur = row[idx_duration] if idx_duration is not None else durations_by_id.get(vid)

            posts.append({
                'platform': 'YouTube',
                'title': str(row[idx_title]).strip(),
                'permalink': u_str,
                'views': int(row[idx_views] or 0) if idx_views is not None else 0,
                'likes': int(row[idx_likes] or 0) if idx_likes is not None else 0,
                'comments': int(row[idx_comments] or 0) if idx_comments is not None else 0,
                'shares': int(row[idx_shares] or 0) if idx_shares is not None else csv_shares_by_id.get(vid, 0),
                'date': date_val,
                'asof': asof,
                'duration_sec': dur,
                'post_id': u_str or str(row[idx_title]),
                'post_type': 'video',
            })
    return posts


def load_youtube_from_studio_csv_only():
    """Fallback when only YouTube Studio's Table data.csv is present (no combined xlsx
    yet). Gives views/shares/duration but NOT likes/comments (0), and has NO date field
    usable as 'Data as of', so asof is left as None rather than guessed."""
    csv_files = glob.glob(os.path.join(RAW_DIR, 'youtube', '*.csv'))
    posts = []
    if not csv_files:
        print('WARNING: no YouTube data found in data/raw/youtube/ (no .xlsx or Table data.csv). '
              'YouTube will be missing from the dashboard.')
        return posts

    print('WARNING: no combined .xlsx found for YouTube — falling back to Table data.csv only. '
          'Likes and comments will show as 0 for YouTube until you re-run yt_pull.py / '
          'yt_merge_shares.py and drop the .xlsx into data/raw/youtube/.')

    with open(csv_files[0], newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row in reader:
            title = row.get('Video title') or row.get('Content') or ''
            if not title or row.get('Content', '').strip() == 'Total':
                continue
            duration = row.get('Duration')
            try:
                duration_sec = int(float(duration)) if duration else None
            except ValueError:
                duration_sec = None
            posts.append({
                'platform': 'YouTube',
                'title': title.strip(),
                'permalink': row.get('Video URL', ''),
                'views': int(float(row.get('Views', 0) or 0)),
                'likes': 0,
                'comments': 0,
                'shares': int(float(row.get('Shares', 0) or 0)) if 'Shares' in row else 0,
                'date': (row.get('Video publish time') or '')[:10] or None,
                'asof': None,
                'duration_sec': duration_sec,
                'post_id': row.get('Video URL') or row.get('Content') or title,
                'post_type': 'video',
            })
    return posts


def _meta_posts(platform, subdir, video_type_values):
    posts = []
    for path in glob.glob(os.path.join(RAW_DIR, subdir, '*.csv')):
        asof = _asof_from_filename(path)
        with open(path, newline='', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                ptype = (row.get('Post type') or row.get('Media type') or '').strip()
                if video_type_values and ptype not in video_type_values:
                    continue
                permalink = row.get('Permalink') or row.get('Post ID') or ''
                if permalink in config.MANUAL_EXCLUDE_PERMALINKS:
                    continue
                title = row.get('Title') or row.get('Description') or row.get('Caption') or ''
                pid = row.get('Post ID') or permalink
                posts.append({
                    'platform': platform,
                    'title': title,
                    'permalink': permalink,
                    'views': int(float(row.get('Views') or row.get('Reach') or 0)),
                    'likes': int(float(row.get('Likes') or row.get('Reactions') or 0)),
                    'comments': int(float(row.get('Comments') or 0)),
                    'shares': int(float(row.get('Shares') or 0)),
                    'date': (row.get('Publish time') or row.get('Date') or '')[:10] or None,
                    'asof': asof,
                    'post_id': pid,
                    'post_type': ptype,
                })
    return posts


def load_facebook():
    return _meta_posts('Facebook', 'facebook', {'Videos', 'Reels'})


def load_instagram():
    return _meta_posts('Instagram', 'instagram', {'IG reel', 'Video'})


def load_tiktok():
    posts = []
    for path in glob.glob(os.path.join(RAW_DIR, 'tiktok', '*.csv')):
        with open(path, newline='', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                title = row.get('Video title') or row.get('Content') or ''
                if not title:
                    continue
                posts.append({
                    'platform': 'TikTok',
                    'title': title,
                    'permalink': row.get('Video link') or row.get('Link', ''),
                    'views': int(float(row.get('Total views') or row.get('Views') or 0)),
                    'likes': int(float(row.get('Total likes') or row.get('Likes') or 0)),
                    'comments': int(float(row.get('Total comments') or row.get('Comments') or 0)),
                    'shares': int(float(row.get('Total shares') or row.get('Shares') or 0)),
                    'date': row.get('Post time') or None,
                    'asof': row.get('Time'),  # literal string from the CSV, e.g. "September 28"
                    'post_id': row.get('Video link') or title,
                    'post_type': 'video',
                })
    return posts


def _dedupe_by_pid(posts):
    """Facebook/Instagram exports can overlap in date range; keep the most-recent-asof
    version of each Post ID."""
    best = {}
    for p in posts:
        pid = p.get('post_id')
        if pid not in best or (p.get('asof') or '') > (best[pid].get('asof') or ''):
            best[pid] = p
    return list(best.values())


def main():
    all_posts = []
    all_posts += load_youtube()
    all_posts += _dedupe_by_pid(load_facebook())
    all_posts += _dedupe_by_pid(load_instagram())
    all_posts += load_tiktok()

    for p in all_posts:
        fix = config.MANUAL_TITLE_TEXT_FIXES.get(p['title'])
        if fix:
            p['title'] = fix

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(all_posts, f, ensure_ascii=False, indent=2)
    print(f'Wrote {len(all_posts)} posts -> {OUT_PATH}')


if __name__ == '__main__':
    main()

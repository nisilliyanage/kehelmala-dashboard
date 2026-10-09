"""
yt_merge_shares.py - add Shares from a YouTube Studio CSV export into youtube_export.xlsx

Run:
  python yt_merge_shares.py "Table data.csv"

Reads:  youtube_export.xlsx   (made by yt_export.py)
Writes: youtube_export_with_shares.xlsx   (your original is left untouched)
"""
import csv
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font

CSV_PATH = sys.argv[1] if len(sys.argv) > 1 else "Table data.csv"
SRC = "youtube_export.xlsx"
OUT = "youtube_export_with_shares.xlsx"

# --- read the Studio CSV: video ID -> shares ---
with open(CSV_PATH, newline="", encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    id_col = next((h for h in headers
                   if h.strip().lower() in ("content", "video id", "video_id")), None)
    share_col = next((h for h in headers if "share" in h.lower()), None)
    if not id_col or not share_col:
        sys.exit(f"Couldn't find the video ID / Shares columns.\nHeaders in the CSV: {headers}")
    shares = {}
    for row in reader:
        vid = (row.get(id_col) or "").strip()
        val = (row.get(share_col) or "").strip().replace(",", "")
        if vid and val:
            try:
                shares[vid] = int(float(val))
            except ValueError:
                pass

# --- write them into the Metrics sheet ---
wb = load_workbook(SRC)
videos, metrics = wb["Videos"], wb["Metrics"]
blue = Font(name="Arial", size=10, color="0000FF")
matched, missing = 0, []

for r in range(2, videos.max_row + 1):
    url = videos.cell(row=r, column=5).value or ""
    vid = url.split("v=")[-1]
    if vid in shares:
        cell = metrics.cell(row=r, column=6)
        cell.value = shares[vid]
        cell.font = blue
        cell.number_format = "#,##0"
        metrics.cell(row=r, column=11).value = "Shares from YouTube Studio export"
        matched += 1
    else:
        missing.append(videos.cell(row=r, column=1).value)

wb.save(OUT)
print(f"Matched shares for {matched} of {videos.max_row - 1} videos -> {OUT}")
if missing:
    print(f"No shares found for {len(missing)} video(s), e.g.: {missing[:3]}")

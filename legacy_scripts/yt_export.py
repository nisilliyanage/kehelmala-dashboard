"""
yt_export.py - turn kehelmala.db (from yt_pull.py) into youtube_export.xlsx

Run:  python yt_export.py
Needs: pip install openpyxl

Makes two sheets in the same layout as the master workbook:
  Videos  - one row per video (Genre left blank for you to fill in)
  Metrics - one row per video, platform = YouTube, latest snapshot
"""
import sqlite3
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

DB = "kehelmala.db"
OUT = "youtube_export.xlsx"

QUERY = """
SELECT p.title, p.published_at, p.url,
       m.views, m.likes, m.comments, m.captured_at
FROM posts p
JOIN metrics_snapshots m
  ON m.platform = p.platform AND m.post_id = p.post_id
WHERE p.platform = 'youtube'
  AND m.captured_at = (SELECT MAX(captured_at) FROM metrics_snapshots
                       WHERE platform = m.platform AND post_id = m.post_id)
ORDER BY p.published_at DESC
"""

conn = sqlite3.connect(DB)
rows = conn.execute(QUERY).fetchall()
conn.close()

font = Font(name="Arial", size=10)
blue = Font(name="Arial", size=10, color="0000FF")
head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="1F3A5F")
yellow = PatternFill("solid", fgColor="FFFF00")

wb = Workbook()
v = wb.active
v.title = "Videos"
m = wb.create_sheet("Metrics")

v.append(["Video", "Campaign", "Release date", "Genre", "Notes"])
m.append(["Video", "Platform", "Views", "Engagement", "Reactions/Likes",
          "Shares/Reposts", "Comments", "Engagement rate", "Data as of",
          "Source", "Notes"])

for i, (title, published, url, views, likes, comments, captured) in enumerate(rows, start=2):
    release = date.fromisoformat(published[:10])
    as_of = date.fromisoformat(captured[:10])
    v.append([title, None, release, None, url])
    m.append([title, "YouTube", views, f"=SUM(E{i}:G{i})", likes, None, comments,
              f'=IF(C{i}>0,D{i}/C{i},"")', as_of, "YouTube Data API (yt_pull.py)",
              "Shares not available from the public API - add from YouTube Studio"])

for ws, ncols in ((v, 5), (m, 11)):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = head, head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = font
    ws.freeze_panes = "A2"

for r in range(2, v.max_row + 1):
    v.cell(row=r, column=3).number_format = "d mmm yyyy"
    v.cell(row=r, column=4).fill = yellow
    for c in (1, 3):
        v.cell(row=r, column=c).font = blue
for r in range(2, m.max_row + 1):
    for c in (3, 5, 7):
        m.cell(row=r, column=c).number_format = '#,##0'
        m.cell(row=r, column=c).font = blue
    m.cell(row=r, column=4).number_format = '#,##0'
    m.cell(row=r, column=8).number_format = '0.0%'
    m.cell(row=r, column=9).number_format = "d mmm yyyy"
    m.cell(row=r, column=1).font = blue
    m.cell(row=r, column=2).font = blue

for ws, widths in ((v, [50, 24, 14, 16, 45]), (m, [50, 10, 12, 12, 15, 15, 11, 14, 13, 30, 55])):
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w

wb.save(OUT)
print(f"Wrote {len(rows)} videos to {OUT}")
